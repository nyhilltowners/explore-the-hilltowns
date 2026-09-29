"""(a) DEC Mined Land Permits -> 'Mines & Quarries': enrich an MRDS pin within 500 m of the same commodity family, else add a
new pin with real permit dates. (b) Industrial WWTPs -> 'Waste & Contamination' (footprint, or >=1 MGD statewide).
(c) Issued Title V permits -> enrich Title V pins (permit id/dates/URL); add pins only for unmatched facilities with coordinates.
Usage: python3 scripts/dec_permits_ingest.py v969"""
import csv,re,sys,math,collections,openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
csv.field_size_limit(10**9)
ver=sys.argv[1] if len(sys.argv)>1 else 'v969'
BOX=lambda lat,lon: 41.9<=lat<=43.2 and -75.1<=lon<=-73.3
cl=lambda s: ILLEGAL_CHARACTERS_RE.sub('',re.sub(r'\s+',' ',str(s or ''))).strip()
def num(x):
    try: return float(str(x).replace(',',''))
    except: return 0.0
def dist(a,b,c,d): return math.hypot((a-c)*111000,(b-d)*82000)
p='data/points_of_interest.xlsx'; wb=openpyxl.load_workbook(p); stats=collections.Counter()
# ---------- (a) Mined land
ws=wb['Mines & Quarries']; h={c.value:c.column for c in ws[1] if c.value}
STATUS={'A':'active','R':'reclaimed','N':'never permitted','T':'terminated','V':'void','X':'expired','P':'pending','Q':'inquiry','E':'expired, not reclaimed','I':'inactive, not reclaimed'}
def fam(com):
    c=com.lower()
    if 'clay' in c or 'shale' in c: return '🧱'
    if 'sand' in c or 'gravel' in c or 'till' in c or 'topsoil' in c or 'peat' in c: return '🐚'
    if any(k in c for k in ['stone','limestone','dolo','granite','slate','marble','trap','gypsum','salt','talc','wollastonite','garnet']): return '🪨'
    return '⛏️'
mrds=[]
for r in range(2,ws.max_row+1):
    if ws.cell(r,2).value and ws.cell(r,h['Latitude']).value:
        mrds.append((r,float(ws.cell(r,h['Latitude']).value),float(ws.cell(r,h['Longitude']).value),ws.cell(r,h['Glyph']).value))
grid=collections.defaultdict(list)
for m in mrds: grid[(round(m[1],2),round(m[2],2))].append(m)
existing={(str(ws.cell(r,2).value).lower(),round(float(ws.cell(r,h['Latitude']).value or 0),3)) for r in range(2,ws.max_row+1) if ws.cell(r,2).value}
for r in csv.DictReader(open('data/layers/dec_mined_land/Mined_Land_Permits_20260928.csv',encoding='utf-8-sig',errors='replace')):
    lat,lon=num(r['Latitude']),num(r['Longitude'])
    if not lat or not lon: stats['mined: no coords']+=1; continue
    g=fam(r['Commodity']); st=STATUS.get(r['Status'],r['Status'])
    yrs=[d[-4:] if len(d)>=4 else d for d in [cl(r['Initial Permit Date'])[:10],cl(r['Permit Issue Date'])[:10],cl(r['Permit End Date'])[:10]] if d]
    yrs=[re.search(r'\d{4}',y).group(0) for y in [cl(r['Initial Permit Date']),cl(r['Permit Issue Date']),cl(r['Permit End Date'])] if re.search(r'\d{4}',y)]
    span=f"{min(yrs)}–{max(yrs)}" if yrs else 'undated'
    dec=(f"DEC Mined Land permit {cl(r['Mine ID Number'])}: {cl(r['Mine Name']) or '(unnamed)'} — {cl(r['Commodity'])}, status {st}, permitted {span}; permittee {cl(r['Permittee Name'])}; {cl(r['Town'])}, {cl(r['County'])} Co.; "
         f"acres: life-of-mine {num(r['Acres Permitted for Life Of Mine']):g}, affected {num(r['Acres Affected']):g}, reclaimed {num(r['Acres Reclaimed']):g}; reclamation {cl(r['Reclamation Type']) or '—'}; underground {cl(r['Underground Mine'])}; last inspected {cl(r['Date Of Last Inspection'])[:10] or '—'}.")
    # match
    best=None
    for dx in (-0.01,0,0.01):
        for dy in (-0.01,0,0.01):
            for m in grid[(round(lat+dx,2),round(lon+dy,2))]:
                d=dist(lat,lon,m[1],m[2])
                if d<500 and m[3]==g and (best is None or d<best[0]): best=(d,m)
    if best:
        row=best[1][0]; n=ws.cell(row,h['Notes']).value or ''
        if cl(r['Mine ID Number']) in n: stats['mined: already']+=1; continue
        ws.cell(row,h['Notes']).value=n+f" 2026-09-28 ({ver}): matched {best[0]:.0f} m away — {dec} Coordinate tightened to DEC's {lat:.5f}, {lon:.5f} (MRDS value kept in notes: {best[1][1]:.5f}, {best[1][2]:.5f})."
        ws.cell(row,h['Latitude']).value=round(lat,6); ws.cell(row,h['Longitude']).value=round(lon,6)
        t=(ws.cell(row,h['Tags']).value or '').replace('; Coordinates Needed','')+f"; DEC Mined Land; Permit {st}"
        ws.cell(row,h['Tags']).value=t
        nm=ws.cell(row,2).value
        if 'permitted' not in nm: ws.cell(row,2).value=re.sub(r'\)$',f"; permitted {span})",nm)
        stats['mined: enriched MRDS pin']+=1
    else:
        name=f"{cl(r['Mine Name']) or cl(r['Permittee Name'])+' pit'} (permitted {span}; {st})"
        key=(name.lower(),round(lat,3))
        if key in existing: stats['mined: dup']+=1; continue
        KIND={'🐚':'Sand & Gravel Pit','🧱':'Clay Pit','🪨':'Quarry','⛏️':'Mine'}[g]
        tags='; '.join(['Mines & Quarries',KIND,cl(r['Commodity']),f"Permit {st}",f"{cl(r['County'])} County",cl(r['Town']),'DEC Mined Land']+(['Underground'] if r['Underground Mine']=='Yes' else []))
        row=ws.max_row+1
        for k,v in {'Name':name,'Tags':tags,'Address':f"{cl(r['Town'])}, {cl(r['County'])} County, NY",'Latitude':round(lat,6),'Longitude':round(lon,6),'Glyph':g,'Display':'Yes','Anchor Type':'exact','Notes':f"2026-09-28: added per Laurie ({ver}) from NYSDEC Mined Land Permits: Beginning 1974 (data.ny.gov export 2026-09-28). {dec} No MRDS site within 500 m of the same commodity family."}.items(): ws.cell(row,h[k]).value=v
        existing.add(key); grid[(round(lat,2),round(lon,2))].append((row,lat,lon,g)); stats['mined: new pin']+=1; stats[g]+=1
# ---------- (b)(c) Waste layer
ws=wb['Waste & Contamination']; h={c.value:c.column for c in ws[1] if c.value}
existing={(str(ws.cell(r,2).value).lower(),round(float(ws.cell(r,h['Latitude']).value or 0),3)) for r in range(2,ws.max_row+1) if ws.cell(r,2).value}
for r in csv.DictReader(open('data/layers/dec_waste_extras/Industrial_WWTP_20260928.csv',encoding='utf-8-sig',errors='replace')):
    lat,lon=num(r['Latitude']),num(r['Longitude'])
    if not lat: stats['iwtp: no coords']+=1; continue
    flow=num(r['Average Design Hydraulic Flow']); inbox=BOX(lat,lon); major=flow>=1.0
    show='Yes' if (inbox or major) else 'No'
    name=f"{cl(r['Facility Name']).title()} — industrial wastewater discharge to {cl(r['Ground or Surface']).lower() or '?'} water ({flow:g} MGD design flow)"
    key=(name.lower(),round(lat,3))
    if key in existing: stats['iwtp: dup']+=1; continue
    tags='; '.join(['Waste & Contamination','Industrial Wastewater','SPDES '+cl(r['SPDES Permit Number']),f"Discharge to {cl(r['Ground or Surface'])} Water"]+(['Major Facility','Major Discharger'] if major else [])+['NYSDEC SPDES'])
    notes=(f"2026-09-28: added per Laurie ({ver}) from NYSDEC Industrial Wastewater Treatment Plants (data.ny.gov export 2026-09-28). SPDES {cl(r['SPDES Permit Number'])}; average design hydraulic flow {flow:g} MGD; discharges to {cl(r['Ground or Surface']).lower()} water. Coordinates in the source are rounded to 0.01° (~1 km) — Anchor approximate. "
           +("Display=Yes: inside footprint." if inbox else ("Display=Yes: ≥1 MGD major discharger statewide per Laurie 2026-09-28." if major else "Display=No: outside footprint, <1 MGD — kept in master.")))
    row=ws.max_row+1
    for k,v in {'Name':name,'Tags':tags,'Address':', '.join(x for x in [cl(r['Street']).title(),cl(r['City']).title(),f"NY {cl(r['Zip Code'])[:5]}"] if x and x!='NY '),'Latitude':round(lat,6),'Longitude':round(lon,6),'Glyph':'🌀','Notes':notes,'Display':show,'Anchor Type':'approximate'}.items(): ws.cell(row,h[k]).value=v
    existing.add(key); stats['iwtp: added']+=1; stats['iwtp: visible']+=show=='Yes'
# Title V permits → enrich by DEC ID derived from permit id (digits before '/')
tv={}
for r in range(2,ws.max_row+1):
    n=ws.cell(r,h['Notes']).value or ''
    m=re.search(r'DEC ID (\d{10})',n)
    if m: tv[m.group(1)]=r
for r in csv.DictReader(open('data/layers/dec_waste_extras/Issued_Title_V_Permits_20260928.csv',encoding='utf-8-sig',errors='replace')):
    did=re.sub(r'\D','',r['PERMIT ID'].split('/')[0])
    if did in tv:
        row=tv[did]; n=ws.cell(row,h['Notes']).value or ''
        if r['PERMIT ID'] in n: continue
        ws.cell(row,h['Notes']).value=n+f" 2026-09-28 ({ver}): current Title V permit {cl(r['PERMIT ID'])} issued {cl(r['ISSUE DATE'])}, expires {cl(r['EXPIRATION DATE'])}; permit text: {cl(r['URL TO PERMIT TEXT'])} (NYSDEC Issued Title V Facility Permits export 2026-09-28)."
        ws.cell(row,h['Website']).value=None
        stats['titlev: enriched']+=1
    else:
        m=re.match(r'POINT \((-?[\d.]+) (-?[\d.]+)\)',r['georeference'] or '')
        if not m: stats['titlev: unmatched, no coords']+=1; continue
        lon,lat=float(m.group(1)),float(m.group(2)); inbox=BOX(lat,lon)
        name=f"{cl(r['FACILITY NAME']).title()} — Title V air permit (issued {cl(r['ISSUE DATE'])[-4:]})"
        key=(name.lower(),round(lat,3))
        if key in existing: continue
        row=ws.max_row+1
        for k,v in {'Name':name,'Tags':'Waste & Contamination; Air Emissions; Title V Facility; Major Facility; NYSDEC Title V','Address':', '.join(x for x in [cl(r['FACILITY LOCATION']).title(),cl(r['FACILITY CITY']).title(),f"NY {cl(r['FACILITY ZIP'])[:5]}"] if x and x!='NY '),'Latitude':round(lat,6),'Longitude':round(lon,6),'Glyph':'🏭','Display':'Yes','Anchor Type':'exact','Notes':f"2026-09-28: added per Laurie ({ver}) from NYSDEC Issued Title V Facility Permits (data.ny.gov export 2026-09-28) — not present in the Title V Emissions Inventory. Permit {cl(r['PERMIT ID'])} issued {cl(r['ISSUE DATE'])}, expires {cl(r['EXPIRATION DATE'])}; permit text: {cl(r['URL TO PERMIT TEXT'])}. Display=Yes: Title V = major source by definition."}.items(): ws.cell(row,h[k]).value=v
        existing.add(key); stats['titlev: new pin']+=1
wb.save(p); print(dict(stats))
