#!/usr/bin/env python3
"""events_inspect.py — sanity-check an events drop-in BEFORE building.
  python3 events_inspect.py <path-to-events__NN.xlsx>
Confirms the 4 recurrence columns exist, and that Model A holds
(all rows dated, zero aggregates, zero duplicate name+date pairs).
"""
import openpyxl, sys
from collections import Counter
path = sys.argv[1] if len(sys.argv) > 1 else 'data/events.xlsx'
wb = openpyxl.load_workbook(path, read_only=True)
ws = wb['Events']; hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
H = {h: i for i, h in enumerate(hdr) if h}
def g(r, k):
    i = H.get(k); return r[i] if i is not None and i < len(r) else None
def isnum(v):
    try: float(v); return True
    except: return False
for col in ['Recur Weeks','Recur Days','Recur Except','Recur Time']:
    print(f'  has "{col}": {col in H}')
pub = dated = drec = nocoord = 0; dupe = Counter()
for r in ws.iter_rows(min_row=2, values_only=True):
    nm = g(r, 'Event Name')
    if not nm or not str(nm).strip(): continue
    if str(g(r, 'Display') or '').strip().lower() in ('no','n','false'): continue
    pub += 1
    hd = g(r, 'Start Date') not in (None, '')
    if hd: dated += 1
    if hd and str(g(r, 'Recur Weeks') or '').strip(): drec += 1
    if not (isnum(g(r, 'Latitude')) and isnum(g(r, 'Longitude'))): nocoord += 1
    if hd: dupe[(str(nm).strip(), str(g(r, 'Start Date'))[:10])] += 1
wb.close()
print(f'published {pub} | dated {dated} | aggregate {pub-dated} | recurring {drec} | '
      f'coordless {nocoord} | dupes {len([1 for v in dupe.values() if v>1])}')
print('Model A OK' if pub == dated and not any(v>1 for v in dupe.values()) else 'CHECK: aggregates or dupes present!')
