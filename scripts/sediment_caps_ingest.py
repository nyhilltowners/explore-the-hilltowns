"""Enrich 'Waste & Contamination' remediation pins with DEC Sediment Caps (v974). Usage: python3 scripts/sediment_caps_ingest.py v974
The export is polygon attributes only (no coordinates) — each row carries the DEC program number, so it is joined to the
remediation pin ingested in v968 by that code. Matched pins gain tags, a Notes paragraph and Display=Yes (contaminated-sediment
sites are legacy/major sites statewide). Unmatched rows are appended hidden with 'Coordinates Needed'.
Shape__Area/Shape__Length are in ground metres (checked: Scajaquada '1,600-foot section' strip has perimeter 1,020 m ≈ 2×488 m;
Onondaga polygon 1.91 km² = 472 acres vs DEC's '~480 acres') — NOT Web Mercator units, so no cos(lat) correction is applied."""
import csv,re,sys,math,collections,openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
ver=sys.argv[1] if len(sys.argv)>1 else 'v974'; SRC='data/layers/dec_waste_extras/Sediment_Caps_20260928.csv'
cl=lambda s: ILLEGAL_CHARACTERS_RE.sub('',re.sub(r'\s+',' ',str(s or ''))).strip()
p='data/points_of_interest.xlsx'; wb=openpyxl.load_workbook(p); ws=wb['Waste & Contamination']; h={c.value:c.column for c in ws[1] if c.value}
rows=list(csv.DictReader(open(SRC,encoding='utf-8-sig')))
bycode=collections.defaultdict(list)
for r in rows: bycode[cl(r['PROGNO'])].append(r)
# index remediation pins by program number (Notes carry "program number XXXX")
idx={}
for i in range(2,ws.max_row+1):
    n=ws.cell(i,h['Notes']).value
    if n and 'Remediation Sites' in str(n):
        m=re.search(r'program number ([A-Z0-9]+)',str(n))
        if m: idx.setdefault(m.group(1),i)
stats=collections.Counter()
def para(r,lat):
    typ=cl(r['TYPE']); short=cl(r['SHORTDESCR']); desc=cl(r['CAPDESCRIP']); cont=cl(r['CONTAMINAN']).replace(' | ','; ')
    L=float(r['Shape__Length'] or 0); A=float(r['Shape__Area'] or 0)
    acres=A/4046.856; km=L/1000
    return (f"Sediment cap (DEC Sediment Caps layer, export 2026-09-28): status {typ}; {short}; contaminants in sediment: {cont}. "
            f"Cap: {desc.rstrip('.')}. DEC polygon perimeter {km:.2f} km, area {A:,.0f} m² (≈{acres:.1f} acres). "
            f"Documents: {cl(r['DOCUMENTS'])} · Registry: {cl(r['REGISTRY'])}.")
for code,R in bycode.items():
    i=idx.get(code)
    if i:
        lat=float(ws.cell(i,h['Latitude']).value)
        tags=ws.cell(i,h['Tags']).value or ''
        add=['Contaminated Sediment','Sediment Cap: '+cl(R[0]['TYPE'])]
        for t in add:
            if t not in tags: tags+='; '+t
        ws.cell(i,h['Tags']).value=tags
        paras=' '.join(para(r,lat) for r in R)
        ws.cell(i,h['Notes']).value=cl(str(ws.cell(i,h['Notes']).value or ''))+f" 2026-09-28 ({ver}, per Laurie): {paras}"
        if ws.cell(i,h['Display']).value!='Yes':
            ws.cell(i,h['Display']).value='Yes'; ws.cell(i,h['Notes']).value+=f" Display flipped to Yes in {ver}: contaminated-sediment site (legacy/major, visible statewide)."; stats['made visible']+=1
        stats['enriched']+=1
    else:
        r=R[0]; row=ws.max_row+1
        vals={'Name':cl(r['SITENAME'])+' — contaminated sediment, '+cl(r['TYPE']).lower(),'Tags':'Waste & Contamination; Remediation Site; Contaminated Sediment; Sediment Cap: '+cl(r['TYPE'])+'; Legacy Site; NYSDEC Remediation; Coordinates Needed',
              'Glyph':'☣️','Display':'No','Anchor Type':'none','Notes':f"2026-09-28: added per Laurie ({ver}) from DEC Sediment Caps export — no matching remediation pin (program number {code}) and the export has no coordinates. "+para(r,0)}
        for k,v in vals.items(): ws.cell(row,h[k]).value=v
        stats['added hidden']+=1
wb.save(p); print(dict(stats), 'codes',len(bycode),'rows',len(rows))
