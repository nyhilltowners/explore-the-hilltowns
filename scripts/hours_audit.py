#!/usr/bin/env python3
"""hours_audit.py — flag hour-integrity problems across POIs.
Two checks:
 1. IDENTICAL hour-sets shared by 3+ differently-named places (copy-paste
    contamination suspects — how the Greenville Pantry / Gathering Hope bug
    was found). Bus stops / chains / dual-listings are expected; eyeball the rest.
 2. Event/venue-type places (hall, lodge, commons, event space, sanctuary,
    community center, grange) carrying posted clock hours — often stray hours
    on an events-only venue (Wisdom Roots, Hilltown Commons, Conkling Hall,
    Unbridled were all this). Some are legit (beer halls, town halls) — verify.
Run from the atlas root.
"""
import openpyxl, re
from collections import defaultdict
DAYS = ['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday','Sunday']
VENUE = re.compile(r'\b(event space|event venue|venue|community center|community hall|grange|american legion|vfw|fire (hall|house|department|company)|banquet|hall|pavilion|amphitheater|fairgrounds|commons|lodge|meeting (house|hall)|sanctuary)\b', re.I)

def has_clock(joined): return re.search(r'\d\s*(am|pm|:)', joined.lower())

wb = openpyxl.load_workbook('data/points_of_interest.xlsx', read_only=True)
sig = defaultdict(list); venue_hits = []
for sn in wb.sheetnames:
    if sn in ('EXAMPLE — not published',): continue
    ws = wb[sn]
    try: hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    except StopIteration: continue
    H = {h: i for i, h in enumerate(hdr) if h}
    if 'Name' not in H: continue
    for r in ws.iter_rows(min_row=2, values_only=True):
        nm = r[H['Name']]
        if not nm: continue
        if str(r[H['Display']] or '').strip().lower() in ('no','n','false'): continue
        hrs = tuple((r[H[d]] if d in H else None) for d in DAYS)
        joined = ' '.join(str(x or '') for x in hrs)
        if not has_clock(joined): continue
        if 'appoint' in joined.lower() or '24' in joined or 'dawn' in joined.lower(): continue
        sig[hrs].append((sn, str(nm)))
        blob = str(nm) + ' ' + str(r[H.get('Tags',0)] or '')
        if VENUE.search(blob): venue_hits.append((sn, str(nm), joined))
wb.close()

print('=== identical hour-sets shared by 3+ places (copy-paste suspects) ===')
for hrs, places in sorted(sig.items(), key=lambda kv: -len(kv[1])):
    names = set(p[1] for p in places)
    if len(names) >= 3:
        compact = ', '.join(f'{d[:2]} {v}' for d, v in zip(DAYS, hrs) if v and str(v).strip().lower() != 'closed')
        print(f'  [{len(names)}] {compact[:55]}')
        for sn, nm in sorted(set(places)): print(f'       [{sn[:12]:12s}] {nm[:40]}')
print()
print('=== venue-type places WITH posted hours (verify events-only) ===')
for sn, nm, j in venue_hits: print(f'  [{sn[:12]:12s}] {nm[:40]}')
