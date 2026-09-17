#!/usr/bin/env python3
"""tag_merge.py — apply tag deletes + consolidations across POIs + events.
Edit the DELETE set and REMAP dict below, then run from the atlas root.
De-duplicates within each cell after remapping. Always work from tag NAMES
(one-off row numbers shift between regenerations of the audit list).

  python3 tag_merge.py     # applies, prints counts

Canonical style for REMAP targets: Title Case, singular nouns, hyphenated
compounds (Dog-Friendly, Gluten-Free Options, Family-Owned).
"""
import openpyxl, re

# --- EDIT THESE ---
DELETE = set()   # e.g. {'Academic', 'August', 'Delmar'}
REMAP = {}       # e.g. {'Artisan': 'Artists & Craftspeople', 'Volunteers': 'Volunteer'}
# ------------------

def split_tags(s):
    if not s: return []
    return [t.strip() for t in re.split(r'[;,]', str(s)) if t.strip()]

def run():
    deleted = remapped = cells = 0
    for path in ['data/points_of_interest.xlsx', 'data/events.xlsx']:
        wb = openpyxl.load_workbook(path)
        for sn in wb.sheetnames:
            if sn == 'EXAMPLE — not published': continue
            ws = wb[sn]
            try: hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            except StopIteration: continue
            H = {h: i+1 for i, h in enumerate(hdr) if h}
            if 'Tags' not in H: continue
            tcol = H['Tags']
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=tcol)
                tags = split_tags(cell.value)
                if not tags: continue
                new, seen, lc = [], set(), False
                for t in tags:
                    if t in DELETE: deleted += 1; lc = True; continue
                    nt = REMAP.get(t, t)
                    if nt != t: remapped += 1; lc = True
                    if nt.lower() not in seen:
                        seen.add(nt.lower()); new.append(nt)
                if lc or len(new) != len(tags):
                    cell.value = '; '.join(new) if new else None
                    cells += 1
        wb.save(path)
    print(f'deleted: {deleted} | remapped: {remapped} | cells changed: {cells}')

if __name__ == '__main__':
    run()
