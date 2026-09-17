#!/usr/bin/env python3
"""tag_audit.py — scan every tag across POIs + events.
Outputs: total unique tags, frequency distribution, variant clusters
(same normalized key = likely singular/plural/case/hyphen dupes), and a
numbered one-off list. Run from the atlas root (where data/ lives).

  python3 tag_audit.py                 # prints summary + writes one_off_tags_<date>.md
"""
import openpyxl, re, collections, datetime, sys
from collections import Counter

def split_tags(s):
    if not s: return []
    return [t.strip() for t in re.split(r'[;,]', str(s)) if t.strip()]

def collect():
    tot = Counter()
    for path in ['data/points_of_interest.xlsx', 'data/events.xlsx']:
        wb = openpyxl.load_workbook(path, read_only=True)
        for sn in wb.sheetnames:
            if sn == 'EXAMPLE — not published': continue
            ws = wb[sn]
            try: hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            except StopIteration: continue
            H = {h: i for i, h in enumerate(hdr) if h}
            if 'Tags' not in H: continue
            for r in ws.iter_rows(min_row=2, values_only=True):
                for t in split_tags(r[H['Tags']]): tot[t] += 1
        wb.close()
    return tot

def nk(t):  # normalize key: lowercase, strip non-alnum, crude singular
    return re.sub(r'[^a-z0-9]', '', t.lower().strip().rstrip('s'))

if __name__ == '__main__':
    tot = collect()
    once = sum(1 for c in tot.values() if c == 1)
    print(f'unique tags: {len(tot)} | used once: {once} ({100*once//len(tot)}%) | 3+: {sum(1 for c in tot.values() if c>=3)}')
    nm = collections.defaultdict(list)
    for t in set(tot): nm[nk(t)].append(t)
    dupes = sorted([(k, sorted(set(v))) for k, v in nm.items() if len(set(v)) > 1])
    print(f'variant clusters (same normalized key): {len(dupes)}')
    for k, vs in dupes:
        print('   ' + ' / '.join(f'{v} ({tot[v]})' for v in vs))
    # write one-off list
    ones = sorted([t for t, c in tot.items() if c == 1], key=str.lower)
    stamp = datetime.date.today().isoformat()
    out = [f'# One-off tags (used exactly once) — {len(ones)} total', '',
           f'Generated {stamp}.', '', '| # | Tag |', '|---|---|']
    for i, t in enumerate(ones, 1): out.append(f'| {i} | {t} |')
    fn = f'one_off_tags_{stamp}.md'
    open(fn, 'w').write('\n'.join(out))
    print(f'wrote {fn}')
