#!/usr/bin/env python3
"""Split the 'Waste & Contamination' sheet of data/points_of_interest.xlsx into one sheet per kind of site (v997).

Each row goes to a sheet by its SECOND tag (the type tag every ingest script writes after the category tag), and its
first tag is rewritten to the new sheet name, e.g.
    'Waste & Contamination; Hazardous Spill Site; …'  →  sheet 'Spills', tags 'Spills; Hazardous Spill Site; …'

Idempotent: if the old sheet is gone there is nothing to do. The older one-shot ingest scripts (bulk_storage_ingest.py,
remediation_cso_ingest.py, waste_extras_ingest.py, …) still write to 'Waste & Contamination'; if one is ever re-run, run
this afterwards and its rows are routed into the new sheets (appended below what is already there).

Usage: python3 scripts/split_waste.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "points_of_interest.xlsx"
OLD = "Waste & Contamination"

# new sheet  ←  type tags (second tag)
GROUPS = {
    "Spills": ["Hazardous Spill Site"],
    "Fuel & Chemical Storage": ["Bulk Storage"],
    "Cleanup Sites": ["Remediation Site", "Waterbody Segment"],
    "Orphaned Wells": ["Orphaned Well"],
    "Landfills & Waste Handling": ["Landfill", "Transfer Station", "Recycling & Processing", "Vehicle Dismantling",
                                   "Waste Tires", "Legacy Site", "Waste Facility", "Waste Combustion",
                                   "Sludge, Biosolids & Hazardous Handling"],
    "Sewage & Overflows": ["Sewage Treatment Plant", "Combined Sewer Overflow"],
    "Industrial Discharges": ["Industrial Wastewater", "Industrial Stormwater"],
    "Air Emissions": ["Air Emissions"],
}
TYPE_TO_SHEET = {t: s for s, ts in GROUPS.items() for t in ts}


def main() -> int:
    wb = load_workbook(XLSX)
    if OLD not in wb.sheetnames:
        print(f"no '{OLD}' sheet — already split")
        return 0
    src = wb[OLD]
    hdr = [c.value for c in src[1]]
    ti = hdr.index("Tags")
    pos = wb.sheetnames.index(OLD)
    buckets = {s: [] for s in GROUPS}
    unknown = {}
    for row in src.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        tags = [t.strip() for t in str(row[ti] or "").split(";")]
        kind = tags[1] if len(tags) > 1 else ""
        sheet = TYPE_TO_SHEET.get(kind)
        if not sheet:
            unknown[kind] = unknown.get(kind, 0) + 1
            continue
        row = list(row)
        row[ti] = "; ".join([sheet] + tags[1:])
        buckets[sheet].append(row)
    if unknown:
        print(f"STOP: type tags with no sheet assigned, add them to GROUPS: {unknown}")
        return 1
    for i, (name, rows) in enumerate(buckets.items()):
        if name in wb.sheetnames:
            ws = wb[name]
        else:
            ws = wb.copy_worksheet(src)          # header styles, widths, freeze panes
            ws.title = name
            ws.delete_rows(2, ws.max_row)
            wb.move_sheet(ws, offset=(pos + i) - wb.sheetnames.index(name))
        for r in rows:
            ws.append(r)
        print(f"  {name:28s} +{len(rows):6d} rows ({sum(1 for r in rows if r[hdr.index('Display')] == 'Yes')} displayed)")
    del wb[OLD]
    wb.save(XLSX)
    print(f"'{OLD}' removed; {sum(len(r) for r in buckets.values())} rows now in {len(GROUPS)} sheets")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
